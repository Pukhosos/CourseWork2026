from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from pathlib import Path
from re import compile, Match, Pattern  # noqa: A004
from typing import Final, overload, Self

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet

TABLE_SHEET_NAME: Final[str] = 'TableS2'
HEADER_ROW_NUMBER: Final[int] = 6
FIRST_DATA_ROW: Final[int] = HEADER_ROW_NUMBER + 1
EXPECTED_COLUMN_COUNT: Final[int] = 17
REPRESENTATIVE_MARKER: Final[str] = 'YES'
ASSEMBLY_ACCESSION_PATTERN: Final[Pattern[str]] = compile(r'^(GC[AF]_\d+\.\d+)')


class TypeConverter:
    @staticmethod
    def require_str(value: object, field_name: str) -> str:
        """Convert a required spreadsheet value to str."""

        if value is None:
            raise ValueError(f'Missing required field {field_name!r}.')
        result: str = str(value).strip()
        if len(result) == 0:
            raise ValueError(f'Empty required field {field_name!r}.')
        return result

    @staticmethod
    def optional_str(value: object) -> str | None:
        """Convert an optional spreadsheet value to str."""

        if value is None:
            return None
        result: str = str(value).strip()
        return result if len(result) > 0 else None

    @staticmethod
    def require_int(value: object, field_name: str) -> int:
        """Convert a required spreadsheet value to int."""

        if isinstance(value, bool):
            raise ValueError(f'Invalid integer value for {field_name!r}: {value!r}.')
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        raise ValueError(f'Invalid integer value for {field_name!r}: {value!r}.')

    @staticmethod
    def taxonomy_string(value: object) -> str:
        """Normalize one taxonomy cell while preserving missing-rank markers."""

        if value is None:
            return '-'
        result: str = str(value).strip()
        return result or '-'


@dataclass(frozen=True, slots=True)
class GenomeRecord:
    """A single prokaryotic genome record from Table S2."""

    record_number: int
    is_representative: bool
    gbff: str
    records_in_gbff: int
    total_length_bp: int
    gene_number: int
    species: str
    taxid: int
    taxonomy: tuple[str, ...]

    @classmethod
    def from_row(cls, row: tuple[object, ...]) -> Self:
        if len(row) != EXPECTED_COLUMN_COUNT:
            error_message: str = ': '.join(
                (
                    'Unexpected number of columns',
                    f'expected {EXPECTED_COLUMN_COUNT}, got {len(row)}.',
                )
            )
            raise ValueError(error_message)
        representative_marker: str | None = TypeConverter.optional_str(row[1])
        return cls(
            record_number=TypeConverter.require_int(row[0], '#'),
            is_representative=(representative_marker == REPRESENTATIVE_MARKER),
            gbff=TypeConverter.require_str(row[2], 'gbff'),
            records_in_gbff=TypeConverter.require_int(row[3], 'records_in_gbff'),
            total_length_bp=TypeConverter.require_int(row[4], 'total_length_bp'),
            gene_number=TypeConverter.require_int(row[5], 'gene_number'),
            species=TypeConverter.require_str(row[6], 'Species'),
            taxid=TypeConverter.require_int(row[7], 'taxid'),
            taxonomy=tuple(TypeConverter.taxonomy_string(value) for value in row[8:]),
        )

    @property
    def assembly_accession(self) -> str:
        """Return the NCBI assembly accession without the `_ASM...` suffix."""

        m: Match[str] | None = ASSEMBLY_ACCESSION_PATTERN.match(self.gbff)
        if m is None:
            raise ValueError(f'Cannot extract an accession from {self.gbff!r}.')
        return m.group(1)

    @property
    def domain(self) -> str | None:
        """Return the first taxonomy element *(normally Bacteria or Archaea)*."""

        return self.get_taxonomy_value(0)

    @property
    def superphylum(self) -> str | None:
        """Return the second taxonomy element."""

        return self.get_taxonomy_value(1)

    @property
    def phylum(self) -> str | None:
        """Return the third taxonomy element."""

        return self.get_taxonomy_value(2)

    @property
    def taxonomic_class(self) -> str | None:
        """Return the fourth taxonomy element."""

        return self.get_taxonomy_value(3)

    @property
    def order(self) -> str | None:
        """Return the fifth taxonomy element."""

        return self.get_taxonomy_value(4)

    @property
    def family(self) -> str | None:
        """Return the sixth taxonomy element."""

        return self.get_taxonomy_value(5)

    @property
    def remaining_taxonomy(self) -> tuple[str, ...]:
        """Return taxonomy elements following the nominal Family column."""

        return self.taxonomy[6:]

    def get_taxonomy_value(self, offset: int) -> str | None:
        """Return one taxonomy element *('-' is treated as missing)*."""

        if offset >= len(self.taxonomy):
            return None
        value: str = self.taxonomy[offset]
        return None if value == '-' else value


@dataclass(frozen=True, slots=True)
class GenomeCollection(Sequence[GenomeRecord]):
    """Collection of genome records with common lookup operations."""

    records: tuple[GenomeRecord, ...]

    @property
    def representatives(self) -> tuple[GenomeRecord, ...]:
        """Return records belonging to the representative 275-genome set."""
        return tuple(record for record in self.records if record.is_representative)

    @classmethod
    def from_xlsx(cls, workbook_path: Path) -> Self:
        """Load all Table S2 records from the supplied workbook."""

        if not workbook_path.is_file():
            raise FileNotFoundError(f'Workbook does not exist: {workbook_path}')
        workbook: Workbook = load_workbook(
            filename=workbook_path,
            read_only=True,
            data_only=True,
        )
        if TABLE_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f'Workbook has no {TABLE_SHEET_NAME!r} worksheet.')
        worksheet: Worksheet = workbook[TABLE_SHEET_NAME]
        records: tuple[GenomeRecord, ...] = tuple(
            map(
                GenomeRecord.from_row,
                filter(
                    lambda values: not all(value is None for value in values),
                    map(
                        tuple,
                        worksheet.iter_rows(
                            min_row=FIRST_DATA_ROW,
                            max_col=EXPECTED_COLUMN_COUNT,
                            values_only=True,
                        ),
                    ),
                ),
            )
        )
        collection: Self = cls(records)
        record_numbers: set[int] = {record.record_number for record in collection}
        taxids: list[int] = [record.taxid for record in collection]
        n_repr: int = len(collection.representatives)
        if len(record_numbers) != len(collection):
            raise ValueError('Duplicate record numbers found in Table S2.')
        if n_repr != 275:
            raise ValueError(f'Expected 275 representative genomes, got {n_repr}')
        if len(taxids) == 0:
            raise ValueError('No genome records were loaded.')
        return collection

    def by_taxid(self, taxid: int) -> GenomeRecord | None:
        """Find a genome by its NCBI taxonomy identifier."""

        return next((record for record in self.records if record.taxid == taxid), None)

    def by_assembly_accession(self, assembly_accession: str) -> GenomeRecord | None:
        """Find a genome by an accession such as GCF_000005845.2."""

        return next(
            (
                record
                for record in self.records
                if record.assembly_accession == assembly_accession
            ),
            None,
        )

    @overload
    def __getitem__(self, index: int, /) -> GenomeRecord:
        raise NotImplementedError

    @overload
    def __getitem__(self, index: slice, /) -> tuple[GenomeRecord, ...]:
        raise NotImplementedError

    def __getitem__(
        self,
        index: int | slice,
    ) -> GenomeRecord | tuple[GenomeRecord, ...]:
        return self.records[index]

    def __len__(self) -> int:
        return len(self.records)

    def __iter__(self) -> Iterator[GenomeRecord]:
        return iter(self.records)
